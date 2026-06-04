#!/usr/bin/env python3
"""Приложения Ж и З: руководство программиста и руководство оператора."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
FIRST_INDENT = Cm(1.25)

ROOT = Path(__file__).resolve().parent.parent
OUT_PROGRAMMER = ROOT / "docs" / "Приложение_Ж_Руководство_программиста.docx"
OUT_OPERATOR = ROOT / "docs" / "Приложение_З_Руководство_оператора.docx"


def setup_doc():
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)
    s.left_margin = Cm(3)
    s.right_margin = Cm(1.5)
    return doc


def title_center(doc, text, level=0):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return h


def para(doc, text):
    doc.add_paragraph(text)


def numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
    doc.add_paragraph()


def build_programmer():
    doc = setup_doc()
    title_center(doc, "ПРИЛОЖЕНИЕ Ж", 0)
    title_center(doc, "РУКОВОДСТВО ПРОГРАММИСТА", 0)

    para(
        doc,
        "Настоящий документ предназначен для специалистов, выполняющих установку, "
        "настройку, сопровождение и развитие веб-приложения «3D-визуализация учебного "
        "оборудования» Алтайского промышленно-экономического колледжа (далее — система).",
    )

    doc.add_heading("1. Общие сведения", level=1)
    para(
        doc,
        "Система реализована как клиент-серверное веб-приложение. Серверная часть — "
        "Node.js 20 и фреймворк Express 5; хранение данных — SQLite 3; клиент — HTML5, "
        "CSS3, JavaScript (ES-модули), веб-компонент Google model-viewer для GLB-моделей. "
        "Развёртывание в production рекомендуется через Docker Compose.",
    )
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["Наименование", "Веб-приложение 3D-визуализация учебного оборудования"],
            ["Репозиторий", "https://github.com/may539/Diplom"],
            ["Точка входа сервера", "server.js"],
            ["Порт по умолчанию", "8080 (переменная PORT)"],
            ["Версия package.json", "1.0.0"],
        ],
    )

    doc.add_heading("2. Назначение и функции программы", level=1)
    numbered(
        doc,
        [
            "Публикация каталога специальностей и учебных 3D-экспонатов (GLB) через REST API.",
            "Интерактивный просмотр моделей в браузере (главная страница и view.html).",
            "Администрирование каталога: добавление, изменение, удаление карточек и загрузка GLB.",
            "Генерация QR-кодов со ссылкой на полноэкранный просмотр (/view.html?id=...).",
            "Логирование сканирований QR (таблица scan_logs, файл logs/scan.log).",
            "Справочник FAQ (таблица help_articles, API /api/help-articles).",
            "Защита админ-операций токеном сессии после проверки мастер-пароля (bcrypt).",
        ],
    )

    doc.add_heading("3. Условия выполнения программы", level=1)
    doc.add_heading("3.1. Технические средства", level=2)
    bullets(
        doc,
        [
            "Сервер: Linux (Debian/Ubuntu) или виртуальная машина с Docker Engine 24+.",
            "ОЗУ: не менее 1 ГБ для контейнера; для загрузки крупных GLB — 2 ГБ и выше.",
            "Диск: место под SQLite, журналы и каталог public/models (до 100 МБ на один GLB).",
            "Клиент: ПК, планшет или смартфон с современным браузером (Chrome, Firefox, Safari, Edge).",
        ],
    )
    doc.add_heading("3.2. Программное обеспечение", level=2)
    add_table(
        doc,
        ["Компонент", "Версия / примечание"],
        [
            ["Node.js (локально или в образе)", "20 LTS"],
            ["npm", "из состава Node.js"],
            ["Docker + Docker Compose", "для production"],
            ["SQLite", "встроена (sqlite3 npm)"],
            ["Express, helmet, compression, multer, qrcode, bcrypt", "см. package.json"],
        ],
    )

    doc.add_heading("4. Структура программы", level=1)
    para(doc, "Корневая структура проекта:")
    add_table(
        doc,
        ["Путь", "Назначение"],
        [
            ["server.js", "HTTP-сервер, маршруты API, инициализация БД, middleware"],
            ["index.html, app.js, styles.css", "Главная страница каталога и 3D-просмотра"],
            ["admin.html, admin.js", "Админ-панель загрузки и управления объектами"],
            ["view.html, view.js, view.css", "Полноэкранный просмотр по QR (мобильные)"],
            ["model-viewer-setup.js", "Подключение веб-компонента model-viewer"],
            ["server/algorithms/", "Модули бизнес-логики (дерево специальностей, CRUD, логи)"],
            ["algorithms-client/", "Клиентские алгоритмы дерева и каталога"],
            ["data/equipment.json", "Seed-данные для первичного наполнения SQLite"],
            ["data/equipment.sqlite", "Рабочая база (создаётся при старте)"],
            ["public/models/", "Загруженные файлы .glb"],
            ["logs/scan.log", "Резервный журнал сканирований QR"],
            ["Dockerfile, docker-compose.yml", "Сборка и запуск контейнера"],
        ],
    )

    doc.add_heading("5. Установка и запуск", level=1)
    doc.add_heading("5.1. Запуск в Docker (рекомендуется)", level=2)
    para(doc, "В каталоге проекта выполните:")
    p = doc.add_paragraph()
    r = p.add_run(
        "cp .env.example .env\n"
        "docker compose up -d --build\n"
        "docker compose logs -f"
    )
    r.font.name = "Consolas"
    r.font.size = Pt(10)
    para(
        doc,
        "Сайт доступен по http://<IP-хоста>:8080/ (порт задаётся HOST_PORT в .env). "
        "Именованные тома diplom-data, diplom-models, diplom-logs сохраняют БД, модели и логи "
        "при пересоздании контейнера.",
    )

    doc.add_heading("5.2. Локальный запуск (разработка)", level=2)
    p = doc.add_paragraph()
    r = p.add_run("npm install\nnpm start\nnpm test")
    r.font.name = "Consolas"
    r.font.size = Pt(10)
    para(doc, "Команда npm test проверяет синтаксис JS и соответствие fallback-данных equipment.json.")

    doc.add_heading("5.3. Переменные окружения", level=2)
    add_table(
        doc,
        ["Переменная", "Описание"],
        [
            ["PORT", "Порт HTTP-сервера (по умолчанию 8080)"],
            ["HOST", "Адрес привязки (0.0.0.0)"],
            ["ADMIN_PASSWORD", "Пароль администратора (plain text, для dev)"],
            ["ADMIN_PASSWORD_HASH", "Bcrypt-хэш пароля (приоритет над ADMIN_PASSWORD)"],
            ["PUBLIC_BASE_URL", "Базовый URL для QR (если пусто — из запроса)"],
            ["HOST_PORT", "Проброс порта в docker-compose (файл .env)"],
        ],
    )

    doc.add_heading("6. Описание REST API", level=1)
    para(doc, "Публичные методы (без токена):")
    add_table(
        doc,
        ["Метод", "URL", "Назначение"],
        [
            ["GET", "/api/specialties", "Список специальностей"],
            ["GET", "/api/specialties/tree", "Дерево специальностей (nodes + tree)"],
            ["GET", "/api/equipment?specialtyId=", "Каталог оборудования по специальности"],
            ["GET", "/api/equipment/:id", "Карточка одного экспоната"],
            ["GET", "/api/qr/:id", "PNG QR-кода для экспоната"],
            ["GET", "/api/help-articles", "Список статей FAQ"],
            ["GET", "/catalog/:id", "Лог сканирования + редирект на view.html"],
            ["GET", "/view.html?id=", "Страница просмотра (лог при scan=1)"],
        ],
    )
    para(doc, "Административные методы (заголовок Authorization: Bearer <token>):")
    add_table(
        doc,
        ["Метод", "URL", "Назначение"],
        [
            ["POST", "/api/admin/login", "Выдача токена по паролю { password }"],
            ["GET", "/api/admin/specialties", "Список специальностей для формы"],
            ["GET", "/api/admin/equipment", "Полный каталог для админки"],
            ["POST", "/api/admin/equipment", "Создание карточки (multipart, поле glb)"],
            ["PUT", "/api/admin/equipment/:id", "Обновление карточки"],
            ["DELETE", "/api/admin/equipment/:id", "Удаление + CASCADE scan_logs"],
        ],
    )
    para(
        doc,
        "Токен сессии действует 24 часа. При превышении лимита запросов (100 за 15 мин) "
        "возвращается JSON { error: «Слишком много запросов...» }.",
    )

    doc.add_heading("7. База данных SQLite", level=1)
    add_table(
        doc,
        ["Таблица", "Назначение"],
        [
            ["specialties", "Специальности (id, code, title, parent_id, sort_order)"],
            ["equipment", "Экспонаты: тексты, model, features_json, model_file_hash"],
            ["help_articles", "Статьи FAQ (slug, title, body)"],
            ["scan_logs", "Журнал просмотров по QR (ON DELETE CASCADE)"],
        ],
    )
    para(
        doc,
        "При первом запуске данные импортируются из data/equipment.json, если таблицы пусты. "
        "Миграции схемы выполняет server/algorithms/db-migrate.js (parent_id, sort_order, scan_logs).",
    )

    doc.add_heading("8. Входные и выходные данные", level=1)
    doc.add_heading("8.1. Входные данные", level=2)
    bullets(
        doc,
        [
            "HTTP-запросы пользователей и администратора.",
            "Файлы GLB (multipart, до 100 МБ, только расширение .glb).",
            "Поля формы: specialtyId, title, type, short, description, features, environment, variant.",
            "Seed-файл data/equipment.json при инициализации.",
        ],
    )
    doc.add_heading("8.2. Выходные данные", level=2)
    bullets(
        doc,
        [
            "HTML-страницы и статические ресурсы (CSS, JS, GLB).",
            "JSON-ответы API.",
            "PNG QR-коды (библиотека qrcode).",
            "Записи в equipment.sqlite и logs/scan.log.",
        ],
    )

    doc.add_heading("9. Сообщения программы", level=1)
    add_table(
        doc,
        ["Код / ситуация", "Сообщение / действие"],
        [
            ["401", "Неверный пароль / Требуется авторизация"],
            ["400", "Invalid equipment id / ошибки валидации полей"],
            ["404", "Объект или статья не найдены"],
            ["429", "Слишком много запросов (rate limit)"],
            ["500", "Внутренняя ошибка сервера (лог в консоль контейнера)"],
            ["Multer", "Только файлы .glb"],
        ],
    )

    doc.add_heading("10. Сопровождение и резервное копирование", level=1)
    numbered(
        doc,
        [
            "Резервная копия БД: скопировать том diplom-data или файл data/equipment.sqlite.",
            "Резервная копия моделей: том diplom-models или каталог public/models.",
            "Обновление версии: git pull, docker compose build --no-cache, docker compose up -d.",
            "Сброс данных: docker compose down -v (удаляет тома!).",
            "Смена пароля: задать ADMIN_PASSWORD_HASH в .env и перезапустить контейнер.",
            "Диагностика: docker compose logs -f; проверка GET /api/specialties (healthcheck).",
        ],
    )

    doc.add_heading("11. Расширение функциональности", level=1)
    para(
        doc,
        "Для добавления специальности или экспоната без админки можно отредактировать "
        "data/equipment.json и удалить equipment.sqlite (при следующем старте выполнится повторный seed). "
        "В production предпочтительна админ-панель. Новые API следует регистрировать в server.js; "
        "бизнес-логику выносить в server/algorithms/. Клиентские сценарии — в app.js и algorithms-client/.",
    )

    para(
        doc,
        "Блок-схемы алгоритмов модулей А–Е приведены в приложениях к пояснительной записке; "
        "текст для онлайн-редакторов — в docs/algorithms/Блок-схемы_для_онлайн-редакторов.txt.",
    )

    OUT_PROGRAMMER.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PROGRAMMER)
    return OUT_PROGRAMMER


# --- Оформление руководства оператора (как в примере RadioManagement) ---


def _style_run(run, *, bold=False, italic=False, size=None):
    run.font.name = FONT_NAME
    run.font.size = size or FONT_SIZE
    run.bold = bold
    run.italic = italic
    r = run._element.get_or_add_rPr()
    r.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def apply_operator_doc_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    pf = normal.paragraph_format
    pf.first_line_indent = FIRST_INDENT
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def op_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    _style_run(r, bold=True)


def op_section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    _style_run(r, bold=True)


def op_subsection(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(text)
    _style_run(r, bold=True)


def op_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = FIRST_INDENT
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    _style_run(r)


def op_dash(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = FIRST_INDENT
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run("–\t" + text)
    _style_run(r)


def op_numbered_list(doc, items):
    for i, text in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = FIRST_INDENT
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(f"{i}.\t{text}")
        _style_run(r)


def op_figure(doc, num, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(f"[ Вставьте рисунок {num} ]")
    _style_run(r, italic=True)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.first_line_indent = Cm(0)
    cr = c.add_run(f"Рисунок {num} – {caption}")
    _style_run(cr, italic=True)
    doc.add_paragraph()


def op_table(doc, caption, headers, rows):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(caption)
    _style_run(r, bold=True)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for para in cell.paragraphs:
            para.paragraph_format.first_line_indent = Cm(0)
            for run in para.runs:
                _style_run(run, bold=True, size=Pt(12))
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = val
            for para in cell.paragraphs:
                para.paragraph_format.first_line_indent = Cm(0)
                for run in para.runs:
                    _style_run(run, size=Pt(12))
    doc.add_paragraph()


def build_operator():
    doc = setup_doc()
    apply_operator_doc_styles(doc)

    op_title(doc, "ПРИЛОЖЕНИЕ З")
    op_title(doc, "РУКОВОДСТВО ОПЕРАТОРА")

    op_section(doc, "1. НАЗНАЧЕНИЕ ПРОГРАММЫ")
    op_subsection(doc, "1.1. Область применения")
    op_body(
        doc,
        "Программа «3D-визуализация учебного оборудования» предназначена для демонстрации "
        "учебных экспонатов Алтайского промышленно-экономического колледжа в интерактивном "
        "трёхмерном виде (формат GLB) на интерактивной панели, компьютере в классе или "
        "лаборатории, а также для просмотра тех же моделей на смартфоне студента после "
        "сканирования QR-кода на учебном стенде. Система позволяет просматривать каталог "
        "оборудования по специальностям, вращать и масштабировать 3D-модели, читать справочные "
        "материалы (FAQ), а администратору — добавлять, изменять и удалять карточки экспонатов, "
        "загружать файлы моделей на сервер и формировать QR-коды для печати.",
    )

    op_subsection(doc, "1.2. Обзор данного документа")
    op_body(doc, "Руководство оператора состоит из следующих разделов:")
    op_dash(doc, "Раздел «Назначение программы» содержит общие сведения о программном продукте и данном руководстве;")
    op_dash(doc, "Раздел «Условия выполнения программы» описывает требования к аппаратуре, программному обеспечению и подготовке пользователя;")
    op_dash(doc, "Раздел «Основы работы с программой» знакомит с запуском системы и элементами главной страницы;")
    op_dash(
        doc,
        "Раздел «Выполнение программы» детально описывает все функции: выбор специальности и "
        "оборудования, работу с 3D-просмотрщиком, меню и справку, админ-панель, QR-коды и "
        "мобильный просмотр;",
    )
    op_dash(doc, "Раздел «Сообщения оператору» содержит перечень сообщений, выдаваемых программой, и рекомендации по действиям.")
    op_body(
        doc,
        "Иллюстрации, приведённые в руководстве, являются схематичными и могут отличаться "
        "от реального вида страниц в зависимости от версии программы, браузера и разрешения экрана.",
    )

    op_section(doc, "2. УСЛОВИЯ ВЫПОЛНЕНИЯ ПРОГРАММЫ")
    op_subsection(doc, "2.1. Уровень подготовки пользователей")
    op_body(
        doc,
        "Пользователь должен иметь базовые навыки работы с операционной системой Windows, "
        "Linux или Android/iOS (запуск браузера, работа с адресной строкой, использование мыши, "
        "касаний и клавиатуры). Специальных знаний в области веб-разработки или 3D-моделирования "
        "не требуется. Перед началом работы рекомендуется ознакомиться с разделом «Выполнение программы». "
        "Для операций в админ-панели оператору необходимо знать мастер-пароль, выданный "
        "администратором информационных систем колледжа.",
    )

    op_subsection(doc, "2.2. Необходимые ресурсы")
    op_subsection(doc, "2.2.1. Технические средства")
    op_dash(doc, "Сервер (или виртуальная машина): процессор 1 ГГц и выше, ОЗУ не менее 1 ГБ (рекомендуется 2 ГБ).")
    op_dash(doc, "Свободное дисковое пространство: не менее 500 МБ для программы, базы SQLite и каталога 3D-моделей.")
    op_dash(doc, "Клиентское устройство: монитор или сенсорная панель с разрешением не менее 1024×768; для мобильного просмотра — смартфон с камерой и браузером.")
    op_dash(doc, "Сеть Ethernet или Wi-Fi для доступа клиентов к серверу по HTTP.")
    op_dash(doc, "Клавиатура и мышь (для ПК); для панели — сенсорный ввод.")

    op_subsection(doc, "2.2.2. Программное обеспечение")
    op_dash(doc, "На сервере: Linux (Ubuntu 20.04+, Debian 11+) с Docker Engine и Docker Compose, либо Node.js 20 LTS при локальном запуске.")
    op_dash(doc, "Отдельная СУБД не устанавливается — используется встроенная SQLite.")
    op_dash(doc, "На клиенте: браузер Google Chrome, Mozilla Firefox, Microsoft Edge или Safari актуальной версии с поддержкой WebGL.")

    op_section(doc, "3. ОСНОВЫ РАБОТЫ С ПРОГРАММОЙ")
    op_subsection(doc, "3.1. Запуск системы")
    op_body(
        doc,
        "Запуск выполняет администратор сервера командой docker compose up -d --build "
        "в каталоге проекта (либо npm start при локальной установке). После готовности сервиса "
        "оператор открывает в браузере адрес http://<IP-или-имя-сервера>:8080/ "
        "(порт может отличаться, если задан HOST_PORT в файле .env). На экране появится "
        "главная страница (рисунок 32).",
    )
    op_figure(doc, 32, "Главная страница веб-приложения")

    op_subsection(doc, "3.2. Описание элементов графического интерфейса")
    op_body(doc, "Интерфейс главной страницы состоит из следующих основных частей:")
    op_numbered_list(
        doc,
        [
            "Шапка сайта — название «Учебное оборудование Алтайского Промышленно-Экономического "
            "колледжа» и кнопка меню «бургер» (три горизонтальные полоски) в правом верхнем углу.",
            "Блок приветствия — краткое описание назначения сайта и кнопки быстрого перехода "
            "к разделам «Выбрать специальность» и «Открыть 3D просмотр».",
            "Раздел «Выберите специальность» — карточки специальностей (ОИБ, ПД, ЗЕМ). "
            "В текущей версии активна только «ПД»; остальные помечены «В разработке».",
            "Раздел «Оборудование и 3D модель» — слева прокручиваемый список объектов (кнопки), "
            "справа — просмотрщик model-viewer, заголовок, тип, описание, список особенностей, "
            "кнопки масштаба и «Автопрокрутка».",
            "Меню «бургер» — пункты: «Загрузить модель», «База знаний (FAQ)», «О нас»; "
            "открывается только по нажатию на кнопку.",
            "Подвал — подпись проекта, кнопка GitHub, ссылки «Файлы cookie» и «Политика конфиденциальности».",
        ],
    )

    op_section(doc, "4. ВЫПОЛНЕНИЕ ПРОГРАММЫ")
    op_subsection(doc, "4.1. Просмотр каталога по специальности")
    op_body(
        doc,
        "Каталог оборудования отображается после выбора доступной специальности. "
        "Оператор может просматривать список экспонатов и переключаться между ними.",
    )
    op_subsection(doc, "4.1.1. Выбор специальности ПД")
    op_body(
        doc,
        "Для выбора специальности прокрутите страницу к разделу «Выберите специальность» "
        "(рисунок 33). Левой кнопкой мыши (или касанием) нажмите на карточку «ПД». "
        "Карточки «ОИБ» и «ЗЕМ» недоступны для выбора.",
    )
    op_figure(doc, 33, "Раздел выбора специальности")
    op_body(
        doc,
        "После выбора страница прокрутится к блоку «Оборудование и 3D модель», "
        "в левом списке появятся объекты данной специальности (рисунок 34).",
    )
    op_figure(doc, 34, "Список оборудования после выбора специальности ПД")

    op_subsection(doc, "4.1.2. Выбор объекта оборудования")
    op_body(
        doc,
        "В левом списке нажмите левой кнопкой мыши на название нужного экспоната. "
        "Справа загрузится соответствующая GLB-модель, обновятся текстовые поля карточки "
        "(рисунок 35). Если список пуст, отобразится сообщение «Для этой специальности пока нет объектов» — "
        "обратитесь к администратору для добавления экспонатов.",
    )
    op_figure(doc, 35, "Выбранный экспонат и 3D-просмотрщик")

    op_subsection(doc, "4.2. Управление трёхмерной моделью")
    op_body(
        doc,
        "Просмотрщик model-viewer позволяет интерактивно изучать модель оборудования.",
    )
    op_subsection(doc, "4.2.1. Вращение и масштаб")
    op_body(
        doc,
        "Для вращения модели удерживайте левую кнопку мыши и перемещайте курсор "
        "(на сенсорной панели — перемещайте палец по области модели). "
        "Для изменения масштаба используйте кнопки «−», «100%», «+» под просмотрщиком "
        "(рисунок 36) или жест «щипок» на мобильном устройстве.",
    )
    op_figure(doc, 36, "Управление масштабом и автопрокруткой 3D-модели")

    op_subsection(doc, "4.2.2. Автопрокрутка и источник GLB")
    op_body(
        doc,
        "Кнопка «Автопрокрутка: вкл/выкл» включает или отключает медленное автоматическое "
        "вращение модели. Кнопка «Открыть GLB-источник» открывает файл модели в новой вкладке "
        "браузера (для проверки загрузки файла).",
    )

    op_subsection(doc, "4.3. Работа с меню «бургер»")
    op_body(
        doc,
        "Меню вызывается нажатием на кнопку с тремя полосками в шапке (рисунок 37). "
        "Повторное нажатие или клик вне области меню закрывает список.",
    )
    op_figure(doc, 37, "Кнопка меню «бургер»")
    op_body(doc, "В открывшемся списке (рисунок 38) доступны пункты:")
    op_dash(doc, "«Загрузить модель» — переход к вводу пароля администратора;")
    op_dash(doc, "«База знаний (FAQ)» — справочные статьи;")
    op_dash(doc, "«О нас» — описание проекта.")
    op_figure(doc, 38, "Открытое выпадающее меню")

    op_subsection(doc, "4.3.1. База знаний (FAQ)")
    op_body(
        doc,
        "Выберите пункт «База знаний (FAQ)». Откроется модальное окно со списком статей "
        "(рисунок 39). Нажмите на заголовок статьи, чтобы раскрыть текст. "
        "Закройте окно кнопкой «×» или кликом по затемнённой области.",
    )
    op_figure(doc, 39, "Модальное окно «База знаний (FAQ)»")

    op_subsection(doc, "4.3.2. Раздел «О нас»")
    op_body(
        doc,
        "Пункт «О нас» открывает модальное окно с кратким описанием назначения системы (рисунок 40).",
    )
    op_figure(doc, 40, "Модальное окно «О нас»")

    op_subsection(doc, "4.4. Подвал сайта и юридические сведения")
    op_body(
        doc,
        "В нижней части главной страницы (рисунок 41) расположены кнопка GitHub "
        "(открывает репозиторий проекта в новой вкладке) и ссылки «Файлы cookie» и "
        "«Политика конфиденциальности». Нажатие на ссылку открывает соответствующее "
        "модальное окно (рисунки 42 и 43).",
    )
    op_figure(doc, 41, "Подвал главной страницы")
    op_figure(doc, 42, "Модальное окно «Файлы cookie»")
    op_figure(doc, 43, "Модальное окно «Политика конфиденциальности»")

    op_subsection(doc, "4.5. Админ-панель: вход и добавление экспоната")
    op_body(
        doc,
        "Операции добавления и изменения каталога доступны только авторизованному администратору.",
    )
    op_subsection(doc, "4.5.1. Вход в админ-панель")
    op_body(
        doc,
        "В меню «бургер» выберите «Загрузить модель». В модальном окне (рисунок 44) "
        "введите мастер-пароль и нажмите «Открыть форму». При верном пароле откроется "
        "страница admin.html.",
    )
    op_figure(doc, 44, "Ввод пароля администратора")
    op_body(
        doc,
        "Внимание: не сообщайте пароль студентам. В production используйте пароль, "
        "отличный от значения по умолчанию admin123.",
    )

    op_subsection(doc, "4.5.2. Добавление нового экспоната")
    op_body(
        doc,
        "На странице админ-панели заполните форму (рисунок 45): выберите специальность, "
        "укажите название, тип, краткое и полное описание, особенности (каждая с новой строки). "
        "Укажите URL файла GLB или нажмите «Выбрать файл» и укажите файл с расширением .glb. "
        "Нажмите «Добавить модель и сгенерировать QR».",
    )
    op_figure(doc, 45, "Форма добавления 3D-модели в админ-панели")
    op_body(
        doc,
        "При корректном заполнении полей появится модальное окно с QR-кодом (рисунок 46). "
        "Нажмите «Печать QR» для печати кода и размещения на стенде. Ссылка в QR ведёт на "
        "страницу view.html?id=<идентификатор>.",
    )
    op_figure(doc, 46, "QR-код после добавления модели")

    op_subsection(doc, "4.6. Редактирование и удаление экспоната")
    op_subsection(doc, "4.6.1. Редактирование")
    op_body(
        doc,
        "В блоке «QR-коды и управление объектами» (рисунок 47) найдите карточку объекта "
        "и нажмите «Изменить». Данные подставятся в форму вверху страницы (рисунок 48). "
        "Внесите изменения и нажмите «Сохранить изменения». После успешного сохранения "
        "отобразится обновлённый QR-код.",
    )
    op_figure(doc, 47, "Каталог объектов в админ-панели")
    op_figure(doc, 48, "Форма редактирования экспоната")

    op_subsection(doc, "4.6.2. Удаление")
    op_body(
        doc,
        "Нажмите кнопку «Удалить» в строке объекта. В диалоге подтверждения (рисунок 49) "
        "нажмите «ОК» для удаления или отмените операцию.",
    )
    op_figure(doc, 49, "Подтверждение удаления объекта")
    op_body(
        doc,
        "Внимание: удаление необратимо. Вместе с записью в базе удаляются связанные "
        "записи журнала сканирований; файл GLB на диске сервера также удаляется, "
        "если модель хранилась в каталоге /models/.",
    )

    op_subsection(doc, "4.7. Просмотр по QR-коду на смартфоне")
    op_body(
        doc,
        "Распечатайте QR-код из админ-панели и разместите у учебного стенда. "
        "Студент сканирует код камерой телефона — откроется страница полноэкранного "
        "просмотра (рисунок 50). Управление моделью выполняется жестами. "
        "Факт сканирования фиксируется в журнале сервера.",
    )
    op_figure(doc, 50, "Полноэкранный просмотр на мобильном устройстве (view.html)")

    op_subsection(doc, "4.8. Повторная печать QR")
    op_body(
        doc,
        "Для существующего объекта в админ-панели нажмите «Печать QR» в строке каталога — "
        "откроется окно с актуальным кодом без повторного сохранения карточки.",
    )

    op_section(doc, "5. СООБЩЕНИЯ ОПЕРАТОРУ")
    op_body(
        doc,
        "В ходе работы с программой оператору могут выдаваться сообщения, приведённые в таблице 1.",
    )
    op_table(
        doc,
        "Таблица 1 – Сообщения оператору",
        ["Текст сообщения", "Описание ситуации", "Действия оператора"],
        [
            [
                "«Не удалось загрузить дерево специальностей» / «Не удалось загрузить каталог»",
                "Сервер недоступен или сбой сети при запросе API.",
                "Проверьте работу сервера; обновите страницу (F5); обратитесь к администратору ИТ.",
            ],
            [
                "«Неверный пароль.»",
                "Введён неверный мастер-пароль при входе в админ-панель.",
                "Проверьте раскладку клавиатуры; уточните пароль у администратора.",
            ],
            [
                "«Требуется авторизация»",
                "Истёк срок действия токена админ-сессии (24 ч) или не выполнен вход.",
                "Повторите вход через меню «Загрузить модель».",
            ],
            [
                "«Заполните обязательные поля»",
                "Не заполнены specialtyId, title, type, short или description.",
                "Дополните форму в админ-панели.",
            ],
            [
                "«Укажите URL GLB или загрузите файл»",
                "Не указан ни URL модели, ни файл .glb.",
                "Укажите ссылку или выберите файл на диске.",
            ],
            [
                "«Только файлы .glb.»",
                "Попытка загрузить файл другого формата.",
                "Конвертируйте модель в GLB или укажите URL готового GLB.",
            ],
            [
                "«Для этой специальности пока нет объектов»",
                "В базе нет экспонатов для выбранной специальности.",
                "Добавьте объекты в админ-панели или выберите другую специальность.",
            ],
            [
                "«Объект не найден»",
                "Запись удалена или указан неверный идентификатор.",
                "Обновите каталог; проверьте корректность QR-ссылки.",
            ],
            [
                "«Слишком много запросов. Повторите попытку позже.»",
                "Сработало ограничение частоты запросов (100 за 15 мин).",
                "Подождите 1–2 минуты и повторите действие.",
            ],
            [
                "«Удалить объект из базы?» / диалог подтверждения",
                "Запрос подтверждения удаления в админ-панели.",
                "Нажмите «ОК» для удаления; «Отмена» — для отмены.",
            ],
            [
                "«Изменения сохранены» / «Объект удалён»",
                "Операция выполнена успешно.",
                "Дополнительных действий не требуется.",
            ],
            [
                "Страница не открывается / ошибка соединения",
                "Сервер остановлен или неверный адрес в браузере.",
                "Сообщите администратору ИТ; проверьте docker compose ps и logs.",
            ],
            [
                "Модель не отображается (пустой просмотрщик)",
                "Ошибка загрузки GLB, блокировка CDN model-viewer или слишком большой файл.",
                "Обновите страницу; проверьте интернет; уменьшите размер GLB.",
            ],
        ],
    )

    OUT_OPERATOR.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_OPERATOR)
    return OUT_OPERATOR


def main():
    p = build_programmer()
    o = build_operator()
    print(f"Saved: {p}")
    print(f"Saved: {o}")


if __name__ == "__main__":
    main()
